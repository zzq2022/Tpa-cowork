import sys
import os
import shutil
import pptx
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement

sys.stdout.reconfigure(encoding='utf-8')

SRC_BASELINE = os.path.abspath(r'docs/ppt/TPA CoWork1.0.pptx')
TARGET_PPT11 = os.path.abspath(r'docs/ppt/TPA CoWork1.1.pptx')
TARGET_PPT11_OPT = os.path.abspath(r'docs/ppt/TPA CoWork1.1_optimized.pptx')
PREVIEW_DIR = os.path.abspath(r'docs/ppt/preview_slides')

FONT_ZH = 'Microsoft YaHei'
FONT_EN = 'Segoe UI'

def apply_font_to_run(r, new_size_pt=None, is_bold=None):
    r.font.name = FONT_ZH
    if new_size_pt is not None:
        r.font.size = Pt(new_size_pt)
    if is_bold is not None:
        r.font.bold = is_bold
    
    rPr = r._r.get_or_add_rPr()
    for old_ea in rPr.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}ea'):
        rPr.remove(old_ea)
    ea = OxmlElement('a:ea')
    ea.set('typeface', FONT_ZH)
    rPr.append(ea)

def get_target_font_size(top, left, height, text, current_fz, is_bold):
    if not text:
        return current_fz, is_bold
    
    # 1. Slide Main Title (Top header)
    if top < 60 and (current_fz is None or current_fz > 20):
        return 30.0, True
    
    # 2. Subtitle
    if top < 110 and (current_fz is not None and 13.5 <= current_fz <= 20) and left < 350:
        return 14.5, is_bold
    
    # 3. Major Section Header (一、, 二、, 三、, 阶段一, 实际案例)
    if text.startswith(('一、', '二、', '三、', '四、', '五、', '阶段一', '阶段二', '实际案例')) or (current_fz and 16.0 <= current_fz <= 20.0 and top < 250):
        return 16.0, True
    
    # 4. Symbols / Arrows (keep existing size)
    if text in ['▼', '|', '➔', '→', '•', '✕']:
        return current_fz, is_bold
    
    # 5. Badges / Numbers (01, 02, 1, 2)
    if text.isdigit() and len(text) <= 3 and current_fz and current_fz < 10:
        return 9.5, True
    
    # 6. Fine-tune fractional font sizes into standard integers/half-points
    if current_fz is not None:
        if current_fz > 25:
            return 30.0, is_bold
        elif current_fz > 18:
            return 20.0, is_bold
        elif current_fz > 15:
            return 16.0, is_bold
        elif current_fz > 13.5:
            return 14.0, is_bold
        elif current_fz > 12.0:
            return 13.0, is_bold
        elif current_fz > 10.5:
            return 11.5, is_bold
        elif current_fz > 9.0:
            return 10.0, is_bold
        elif current_fz > 7.5:
            return 9.0, is_bold
        else:
            return 8.5, is_bold
            
    return current_fz, is_bold

def optimize_presentation(src_file, dst_file):
    print(f'Loading baseline from {src_file}...')
    prs = pptx.Presentation(src_file)
    
    modified_count = 0
    for slide_idx, slide in enumerate(prs.slides):
        for shape_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                tf = shape.text_frame
                top = shape.top.pt if (hasattr(shape, 'top') and shape.top) else 0
                left = shape.left.pt if (hasattr(shape, 'left') and shape.left) else 0
                height = shape.height.pt if (hasattr(shape, 'height') and shape.height) else 0
                
                for p_idx, p in enumerate(tf.paragraphs):
                    text = p.text.strip()
                    if not text:
                        continue
                    
                    orig_fz = p.font.size.pt if (p.font and p.font.size) else None
                    p_bold = p.font.bold if p.font else False
                    
                    new_fz, new_bold = get_target_font_size(top, left, height, text, orig_fz, p_bold)
                    
                    # Apply font settings without destroying run colors
                    if p.runs:
                        for r in p.runs:
                            r_fz = r.font.size.pt if (r.font and r.font.size) else orig_fz
                            r_bold = r.font.bold if (r.font and r.font.bold is not None) else new_bold
                            target_r_fz, target_r_bold = get_target_font_size(top, left, height, r.text.strip(), r_fz, r_bold)
                            apply_font_to_run(r, target_r_fz, target_r_bold)
                    else:
                        p.font.name = FONT_ZH
                        if new_fz: p.font.size = Pt(new_fz)
                        if new_bold is not None: p.font.bold = new_bold
                    
                    modified_count += 1

    prs.save(dst_file)
    print(f'Successfully saved optimized PPT to {dst_file} (modified {modified_count} paragraphs)')

if __name__ == '__main__':
    # 1. Restore baseline to PPT1.1
    shutil.copyfile(SRC_BASELINE, TARGET_PPT11)
    
    # 2. Optimize
    optimize_presentation(SRC_BASELINE, TARGET_PPT11)
    shutil.copyfile(TARGET_PPT11, TARGET_PPT11_OPT)
    print('Done optimizing PPT files.')
